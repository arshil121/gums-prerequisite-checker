#!/usr/bin/env python3
"""
Convert data/GUB_Combined_Pre_Course_Data.xlsx -> data/remedial-list.json

The userscript reads the JSON (fast). If the JSON is missing it falls back to
parsing the .xlsx directly, so this script is an optimisation, not a hard
requirement. Run it locally, or let .github/workflows/build-remedial-json.yml
run it automatically whenever the spreadsheet changes.

    python3 tools/build_remedial_json.py
"""
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "GUB_Combined_Pre_Course_Data.xlsx"
OUT = ROOT / "data" / "remedial-list.json"
SHEET = "Combined Students"


def truthy(value):
    return str(value).strip().lower() in {"yes", "y", "1", "true"}


def find_header(rows):
    for index, row in enumerate(rows):
        if any(str(cell).strip().lower() == "student id" for cell in row):
            return index
    raise SystemExit("Could not find a 'Student ID' header cell in the sheet.")


def main():
    if not XLSX.exists():
        raise SystemExit(f"Spreadsheet not found: {XLSX}")

    book = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    sheet = book[SHEET] if SHEET in book.sheetnames else book[book.sheetnames[0]]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]

    header_index = find_header(rows)
    header = [str(c).strip().lower() if c is not None else "" for c in rows[header_index]]

    def col(*names):
        for name in names:
            if name in header:
                return header.index(name)
        return -1

    i_id = col("student id", "student_id")
    i_name = col("student name", "name")
    i_math = col("pre-math", "pre_math")
    i_eng = col("pre-english", "pre_english")
    i_combined = col("pre courses", "pre_courses")

    if i_id == -1:
        raise SystemExit("The sheet has no 'Student ID' column.")

    students = {}
    for row in rows[header_index + 1:]:
        if i_id >= len(row) or row[i_id] in (None, ""):
            continue
        student_id = str(row[i_id]).strip()
        if not student_id.isdigit() or len(student_id) < 6:
            continue

        courses = set()
        if i_math != -1 and truthy(row[i_math]):
            courses.add("Pre-Math")
        if i_eng != -1 and truthy(row[i_eng]):
            courses.add("Pre-English")
        if i_combined != -1 and row[i_combined]:
            text = str(row[i_combined]).lower()
            if "math" in text:
                courses.add("Pre-Math")
            if "english" in text:
                courses.add("Pre-English")
        if not courses:
            continue

        name = str(row[i_name]).strip() if i_name != -1 and row[i_name] else ""
        entry = students.setdefault(student_id, {"id": student_id, "name": name, "_c": set()})
        entry["_c"] |= courses
        if not entry["name"] and name:
            entry["name"] = name

    payload = [
        {"id": e["id"], "name": e["name"], "courses": ", ".join(sorted(e["_c"]))}
        for e in sorted(students.values(), key=lambda e: e["id"])
    ]

    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(payload)} students to {OUT.relative_to(ROOT)} ({OUT.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
